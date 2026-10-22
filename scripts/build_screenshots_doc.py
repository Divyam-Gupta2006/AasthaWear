import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def create_screenshots_doc():
    doc = docx.Document()
    
    # 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
    primary_color = RGBColor(14, 14, 14)
    accent_color = RGBColor(142, 102, 62)
    muted_color = RGBColor(100, 100, 100)
    
    # ---------------------------------------------------------------------------
    # COVER PAGE
    # ---------------------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(80)
    title_p.paragraph_format.space_after = Pt(8)
    title_run = title_p.add_run("आstha — E-commerce Website")
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = primary_color
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(36)
    sub_run = sub_p.add_run("Website Rendered Interface & Visual Proof Document")
    sub_run.font.name = "Georgia"
    sub_run.font.size = Pt(20)
    sub_run.font.italic = True
    sub_run.font.color.rgb = accent_color
    
    divider = doc.add_paragraph()
    divider.paragraph_format.space_after = Pt(60)
    div_run = divider.add_run("―" * 45)
    div_run.font.color.rgb = accent_color
    
    meta_box = doc.add_table(rows=5, cols=2)
    meta_box.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_data = [
        ("Student name:", "[ENTER STUDENT NAME]"),
        ("Roll number:", "[ENTER ROLL NUMBER]"),
        ("Course/subject:", "[ENTER COURSE NAME]"),
        ("Institution:", "[ENTER INSTITUTION]"),
        ("Submission date:", "[ENTER DATE]")
    ]
    for i, (label, val) in enumerate(meta_data):
        row = meta_box.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(6)
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(11)
        r0.font.bold = True
        r0.font.color.rgb = primary_color
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(6)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(11)
        r1.font.color.rgb = muted_color

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # DOCUMENT OVERVIEW
    # ---------------------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r = h1.add_run("Website Rendered Visual Proof")
    r.font.name = "Georgia"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = primary_color
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    p.paragraph_format.line_spacing = 1.2
    p.add_run(
        "This document contains high-fidelity screen captures of the running आstha sacred apparel storefront. "
        "Every capture was recorded from the production build operating at http://localhost:3000 across desktop (1280x840) "
        "and mobile (390x844) viewports. The captures verify complete visual coherence, typography, product details, "
        "interactive galleries, and transactional flows."
    )

    screenshots_dir = os.path.abspath("submission_reports/screenshots")
    
    figures = [
        (
            "Figure 1: Homepage — Hero Campaign & Global Navigation",
            "fig1_homepage_hero.png",
            "Route: / | Viewport: Desktop (1280px)\n"
            "Demonstrates the refined editorial luxury hero banner, announcement bar with rotating delivery guarantees, "
            "desktop navigation links with Collections dropdown, and high-contrast campaign typography."
        ),
        (
            "Figure 2: Shop Catalogue — Tradition Filters & Product Grid",
            "fig2_shop_catalogue.png",
            "Route: /shop | Viewport: Desktop (1280px)\n"
            "Demonstrates the complete catalogue browsing experience, tradition filter chips (Vedic, Priests, Monks), "
            "material and price select filters, sorting control, and consistent 3:4 portrait product cards."
        ),
        (
            "Figure 3: Product Detail — Pitambari Dhoti & Angavastram",
            "fig3_product_pitambari.png",
            "Route: /product/pitambari-dhoti-angavastram | Viewport: Desktop (1280px)\n"
            "Demonstrates the multi-image gallery with thumbnail selection, zoom capability, size selection, package inclusions breakdown, "
            "and technical tabs for verified Ahimsa mulberry silk composition and pure zari border specifications."
        ),
        (
            "Figure 4: Product Detail — Roman Clerical Wool Cassock",
            "fig4_product_cassock.png",
            "Route: /product/roman-clerical-wool-cassock | Viewport: Desktop (1280px)\n"
            "Demonstrates the ecclesiastical clerical cassock detail page featuring 33 hand-stitched cloth buttons, "
            "pure tropical worsted wool specifications, tailored chest sizing (38R-46R), and direct consultation CTA."
        ),
        (
            "Figure 5: Shopping Bag & Considered Order Breakdown",
            "fig5_shopping_cart.png",
            "Route: /cart | Viewport: Desktop (1280px)\n"
            "Demonstrates the shopping bag with live line item details, quantity adjustment controls, item removal, "
            "active voucher code application (ASTHA10 / ATELIER500), and subtotal calculation."
        ),
        (
            "Figure 6: Checkout — Indian Delivery Address & Payment Methods",
            "fig6_checkout_page.png",
            "Route: /checkout | Viewport: Desktop (1280px)\n"
            "Demonstrates the guest checkout workflow with full Indian delivery fields (PIN code, state dropdown, street address), "
            "complimentary insured courier selection, and payment options (UPI, Credit/Debit Card, Net Banking)."
        ),
        (
            "Figure 7: About Us — Sacred Textile Craftsmanship Manifesto",
            "fig7_about_page.png",
            "Route: /about | Viewport: Desktop (1280px)\n"
            "Demonstrates the atelier manifesto detailing sacred textile sanctity, multi-generational Varanasi pit-loom guilds, "
            "and canonical exactitude across Vedic, Buddhist, and Roman liturgical traditions."
        ),
        (
            "Figure 8: Contact Atelier — Sizing & Consultation Desk",
            "fig8_contact_page.png",
            "Route: /contact | Viewport: Desktop (1280px)\n"
            "Demonstrates the interactive consultation form for custom ordination vestment measurements, "
            "parish commissions, direct atelier email, telephone, and physical weaving guild addresses."
        ),
        (
            "Figure 9: Frequently Asked Questions — Canonical Guidance",
            "fig9_faq_page.png",
            "Route: /faq | Viewport: Desktop (1280px)\n"
            "Demonstrates the interactive categorized FAQ accordion covering Ahimsa silk purity, Roman clerical sizing, "
            "diocesan feast day commissions, and insured pan-India delivery."
        ),
        (
            "Figure 10: Mobile Responsive View — Homepage Hero & Navigation",
            "fig10_mobile_homepage.png",
            "Route: / | Viewport: Mobile (390x844 iPhone Viewport)\n"
            "Demonstrates responsive mobile scaling, accessible top bar, hamburger navigation drawer trigger, "
            "and touch-friendly bottom TabBar."
        ),
        (
            "Figure 11: Mobile Responsive View — Shop Catalogue",
            "fig11_mobile_shop.png",
            "Route: /shop | Viewport: Mobile (390x844 iPhone Viewport)\n"
            "Demonstrates mobile product catalogue grid with responsive card scaling, horizontally scrollable tradition chips, "
            "and mobile filter drawer trigger."
        ),
        (
            "Figure 12: Dynamic XML Sitemap — Canonical URL Index",
            "fig12_sitemap_xml.png",
            "Route: /sitemap.xml | Viewport: Desktop (1280px)\n"
            "Demonstrates the rendered XML sitemap containing all 18 canonical static and dynamic product URLs, "
            "calibrated priority ratings, and update frequencies."
        ),
        (
            "Figure 13: Search Crawler Directives — Robots Configuration",
            "fig13_robots_txt.png",
            "Route: /robots.txt | Viewport: Desktop (1280px)\n"
            "Demonstrates the live robots.txt directives properly shielding private utility routes (/cart, /checkout, /account, /wishlist) "
            "while referencing the production XML sitemap."
        )
    ]
    
    for fig_title, img_name, fig_desc in figures:
        img_path = os.path.join(screenshots_dir, img_name)
        if os.path.exists(img_path):
            doc.add_page_break()
            
            p_title = doc.add_paragraph()
            p_title.paragraph_format.space_before = Pt(8)
            p_title.paragraph_format.space_after = Pt(6)
            r = p_title.add_run(fig_title)
            r.font.name = "Georgia"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = primary_color
            
            # Picture width based on aspect ratio
            img_width = Inches(3.2) if "mobile" in img_name else Inches(6.2)
            
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_picture(img_path, width=img_width)
            
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(6)
            p_cap.paragraph_format.space_after = Pt(12)
            p_cap.paragraph_format.line_spacing = 1.15
            r_cap = p_cap.add_run(fig_desc)
            r_cap.font.name = "Arial"
            r_cap.font.size = Pt(9.5)
            r_cap.font.color.rgb = muted_color

    out_path = os.path.abspath("submission_reports/Astha_Website_Screenshots.docx")
    doc.save(out_path)
    print(f"Screenshots document successfully saved to: {out_path}")

if __name__ == "__main__":
    create_screenshots_doc()

